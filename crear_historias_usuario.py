from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = r"C:\laragon\www\proypebieai\Historias_de_Usuario_Finales_PEBIE.docx"

historias = [
    ("HU01", "Iniciar sesión y acceder según el rol", "Usuario del sistema", "Iniciar sesión con mi cuenta y contraseña y visualizar las opciones que correspondan a mi rol.", "Acceder de forma segura solo a las funciones e información necesarias para realizar mi trabajo.", "El control de permisos, el panel por rol y el registro de accesos son reglas internas del sistema."),
    ("HU02", "Gestionar usuarios y asignar roles", "Administrador", "Registrar, actualizar y desactivar usuarios, asignándoles el rol que corresponda.", "Controlar quién utiliza el sistema y qué operaciones puede realizar cada integrante del personal.", "Los roles son administrador, vendedor, encargado de almacén y encargado de despacho."),
    ("HU03", "Consultar indicadores y alertas operativas", "Administrador o encargado de almacén", "Visualizar indicadores de ventas, compras y alertas de stock mínimo según mis funciones.", "Conocer la situación operativa de la empresa y tomar decisiones oportunas.", "El administrador visualiza el resumen general; el encargado de almacén visualiza principalmente alertas de stock mínimo."),
    ("HU04", "Configurar datos de la empresa", "Administrador", "Actualizar razón social, NIT, dirección, teléfono, correo, ciudad, logo y mensaje de pie.", "Mantener correcta la información institucional mostrada en comprobantes y catálogo.", "La configuración representa a Importadora PEBIE."),
    ("HU05", "Gestionar clientes", "Administrador o vendedor", "Registrar, consultar, actualizar y eliminar clientes unipersonales e institucionales.", "Identificar correctamente a cada cliente al realizar ventas o pedidos.", "El registro institucional incluye razón social, NIT y contacto; el unipersonal incluye nombre, apellido y CI."),
    ("HU06", "Gestionar proveedores", "Administrador", "Registrar, consultar, actualizar, desactivar o reactivar proveedores.", "Relacionar las compras con sus proveedores y mantener organizada la información de abastecimiento.", "Se registra empresa, contacto, teléfono, país, calificación y datos complementarios."),
    ("HU07", "Gestionar categorías de artículos", "Administrador", "Registrar, consultar, actualizar y eliminar categorías de artículos.", "Clasificar correctamente los productos y organizar el catálogo comercial.", "No se permite eliminar una categoría que tenga artículos asociados."),
    ("HU08", "Gestionar artículos", "Administrador", "Registrar, consultar, actualizar y desactivar artículos con su categoría, proveedor, precios, costo, imagen y nivel mínimo de stock.", "Mantener actualizado el catálogo comercial y el inventario de artículos para el hogar.", "El vendedor y el encargado de almacén solo pueden consultar los artículos disponibles."),
    ("HU09", "Gestionar ubicaciones físicas", "Administrador o encargado de almacén", "Registrar, consultar, actualizar y eliminar sucursales o almacenes físicos.", "Organizar los espacios donde se resguarda la mercadería.", "No se puede eliminar una ubicación que tenga stock asignado."),
    ("HU10", "Distribuir existencias por ubicación", "Administrador o encargado de almacén", "Asignar y actualizar las cantidades de artículos existentes en cada ubicación física.", "Conocer dónde se encuentra la mercadería y controlar su distribución.", "El sistema impide repartir una cantidad mayor al stock total disponible."),
    ("HU11", "Gestionar compras", "Administrador", "Registrar compras con proveedor, artículos, cantidades y costos; consultar por período o proveedor e imprimir comprobantes individuales o listados.", "Incrementar el inventario y controlar las adquisiciones, cantidades, costos y montos realizados.", "El comprobante se genera automáticamente y el producto puede actualizarse desde un modal de compra."),
    ("HU12", "Registrar mermas de inventario", "Administrador o encargado de almacén", "Registrar artículos dañados, devueltos, perdidos u obsoletos indicando cantidad, motivo y descripción.", "Mantener el stock disponible actualizado después de una baja de inventario.", "No se permite registrar una merma mayor al stock disponible."),
    ("HU13", "Consultar kardex de artículos", "Administrador", "Consultar las entradas, salidas, bajas y saldo de un artículo.", "Conocer la trazabilidad de los movimientos de inventario de cada producto.", "El kardex integra compras, ventas y mermas en orden cronológico."),
    ("HU14", "Gestionar ventas y pedidos", "Administrador o vendedor", "Registrar ventas directas o pedidos indicando cliente, artículos, cantidades, modalidad, medio de pago y promoción; consultar, filtrar, anular e imprimir comprobantes.", "Completar y controlar las operaciones comerciales de la empresa.", "El precio mayorista se aplica automáticamente al alcanzar la cantidad configurada; el vendedor solo visualiza sus propias ventas."),
    ("HU15", "Gestionar despacho de ventas con transporte", "Vendedor o encargado de despacho", "Actualizar el estado de despacho de una venta con promoción de transporte como pendiente, despachado o entregado al cliente.", "Controlar el envío hasta confirmar la entrega efectiva al cliente.", "La venta transportada descuenta stock y se contabiliza como ingreso únicamente al quedar entregada."),
    ("HU16", "Consultar reporte financiero", "Administrador o vendedor", "Consultar los ingresos por ventas y, según mi rol, las compras realizadas en un período.", "Apoyar el control financiero y la revisión de resultados comerciales.", "El administrador visualiza ventas y compras; el vendedor visualiza únicamente el detalle de sus ventas completadas."),
    ("HU17", "Consultar catálogo web y solicitar reserva", "Cliente", "Consultar artículos disponibles y solicitar la reserva de un producto indicando mis datos y cantidad.", "Conocer la oferta de Importadora PEBIE y comunicar mi intención de compra sin realizar un pago en línea.", "El sistema verifica disponibilidad, calcula el precio aplicable, genera un código de reserva y muestra su fecha límite."),
    ("HU18", "Gestionar reservas web", "Personal autorizado", "Consultar y actualizar el estado de las reservas recibidas desde el catálogo web.", "Dar seguimiento a cada solicitud hasta confirmarla, venderla, entregarla o cancelarla.", "Las reservas finalizadas no pueden cambiar nuevamente de estado."),
]

def set_cell_border(cell, color="000000", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tc_pr.append(borders)
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.find(qn(f'w:{side}'))
        if element is None:
            element = OxmlElement(f'w:{side}'); borders.append(element)
        element.set(qn('w:val'), 'single'); element.set(qn('w:sz'), size); element.set(qn('w:color'), color)

def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar'); tc_pr.append(margins)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); margins.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def cell_text(cell, text, bold=False, size=10.5):
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text); run.bold = bold; run.font.name = 'Arial'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial'); run.font.size = Pt(size)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8); section.left_margin = Cm(2); section.right_margin = Cm(2)

styles = doc.styles
styles['Normal'].font.name = 'Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); styles['Normal'].font.size = Pt(10.5)
title = doc.add_paragraph(); title.style = 'Title'; title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Historias de Usuario Finales PEBIE'); run.font.name = 'Arial'; run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); run.font.size = Pt(18); run.font.color.rgb = RGBColor(0,0,0)
intro = doc.add_paragraph('Las siguientes historias de usuario corresponden al alcance funcional actual del Sistema Comercial de Importadora PEBIE. Se incluyen dieciocho historias con valor independiente para sus usuarios, mientras que validaciones, automatismos y consultas auxiliares se expresan como reglas de negocio.')
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; intro.paragraph_format.space_after = Pt(10)

for index, (code, title, como, quiero, para, note) in enumerate(historias, start=1):
    if index > 1:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(4.0); table.columns[1].width = Cm(12.5)
    merged = table.cell(0,0).merge(table.cell(0,1))
    cell_text(merged, f'{code}: {title}', bold=True, size=11)
    labels = ['Como', 'Quiero', 'Para']; values = [como, quiero, para]
    for r, label, value in zip(range(1,4), labels, values):
        cell_text(table.cell(r,0), label, bold=True)
        cell_text(table.cell(r,1), value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell); set_cell_margins(cell)
    note_p = doc.add_paragraph()
    note_p.paragraph_format.space_before = Pt(2); note_p.paragraph_format.space_after = Pt(8)
    note_run = note_p.add_run('Nota. '); note_run.bold = True; note_run.font.name = 'Arial'; note_run.font.size = Pt(9.5)
    text_run = note_p.add_run(note); text_run.font.name = 'Arial'; text_run.font.size = Pt(9.5)
    if index % 2 == 0 and index != len(historias):
        doc.add_page_break()

doc.save(OUTPUT)
print(OUTPUT)

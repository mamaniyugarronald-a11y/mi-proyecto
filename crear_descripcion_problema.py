from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"C:\laragon\www\proypebieai\Descripcion_y_Dominio_del_Problema_Actualizado_PEBIE.docx"

paragraphs = [
    "La Empresa Importadora PEBIE es un establecimiento dedicado a la comercialización de productos y artículos para el hogar. Sus actividades principales comprenden la adquisición de mercadería a proveedores, el almacenamiento y distribución de existencias, el control de inventario, la atención a clientes, las ventas, los pedidos con transporte, el despacho y la presentación de productos mediante un catálogo web.",
    "Antes de la implementación del sistema, gran parte de la información comercial y administrativa se registraba de forma manual mediante cuadernos, hojas sueltas y notas de venta. Esta situación provocaba registros dispersos, demora al consultar información, riesgo de pérdida de datos y diferencias entre las existencias físicas y las cantidades anotadas.",
    "La empresa requiere controlar qué personas pueden acceder al sistema y qué actividades pueden realizar. Para ello se consideran los roles de administrador, vendedor, encargado de almacén y encargado de despacho. El administrador administra usuarios, datos de la empresa y la información general; el vendedor registra ventas y clientes; el encargado de almacén consulta productos, registra mermas y administra ubicaciones físicas; y el encargado de despacho actualiza las ventas que requieren transporte. La ausencia de control centralizado de cuentas y permisos puede comprometer la seguridad y dificultar la asignación de responsabilidades.",
    "En la gestión de clientes, PEBIE atiende a personas particulares e instituciones. Se requiere registrar y mantener actualizados sus datos, distinguiendo clientes unipersonales, con nombre, apellido y cédula de identidad, de clientes institucionales, con razón social, NIT y persona de contacto. Esta información es necesaria para identificar adecuadamente al cliente durante una venta o un pedido.",
    "La gestión de proveedores permite conservar los datos de las empresas que suministran mercadería, incluyendo contacto, teléfono, correo, dirección, país, calificación y observaciones. Tener esta información centralizada facilita relacionar cada compra con su proveedor y consultar los antecedentes de abastecimiento.",
    "Respecto a los artículos, la empresa necesita administrar categorías y productos con información como nombre, descripción, marca, proveedor, imagen, costo de compra, precio por menor, precio por mayor, cantidad mínima para precio mayorista, stock mínimo y estado. La clasificación por categorías y el registro de imágenes permiten organizar tanto la consulta interna como el catálogo público. Sin un registro actualizado, se dificulta conocer la disponibilidad y aplicar los precios correctos.",
    "El inventario requiere un control permanente de entradas, salidas y bajas. Las compras incrementan automáticamente el stock de los productos; las ventas completadas y las mermas lo reducen. El sistema debe alertar sobre productos cuyo stock alcanza el nivel mínimo definido y permitir consultar el kardex por artículo, integrando de forma cronológica los movimientos generados por compras, ventas y mermas. Esto evita depender exclusivamente de revisiones físicas para identificar faltantes o inconsistencias.",
    "La mercadería puede organizarse en almacenes o sucursales físicas. Por ello se requiere registrar ubicaciones y distribuir cantidades de productos entre ellas, sin exceder el stock total disponible. Esta funcionalidad permite conocer dónde se encuentra la mercadería y evita asignar una misma existencia a más de una ubicación de forma inconsistente.",
    "El proceso de compra inicia con la selección del proveedor, los productos y las cantidades adquiridas. Al registrar una compra, el sistema genera un comprobante y actualiza inmediatamente el stock. Además, permite consultar las compras del día o de un período determinado, filtrar por proveedor e imprimir comprobantes individuales o listados del período. De esta manera, las adquisiciones, costos y cantidades quedan concentrados en una sola fuente de información.",
    "La empresa también debe registrar las mermas ocasionadas por daño, devolución, pérdida, defecto u obsolescencia. Cada baja requiere identificar el artículo, cantidad, motivo y descripción, verificando que la cantidad no sea mayor al stock disponible. El registro oportuno de mermas contribuye a que las existencias mostradas por el sistema representen la disponibilidad real.",
    "En la atención comercial, el vendedor registra ventas directas o pedidos mediante un punto de venta. La operación incluye cliente, productos, cantidades, medio de pago, modalidad y promoción. El sistema aplica automáticamente el precio mayorista cuando la cantidad de un producto alcanza el mínimo configurado; además, valida la disponibilidad antes de completar la venta. Las ventas pueden consultarse por fechas, estado, tipo de operación o estado de despacho, y sus comprobantes pueden imprimirse de forma individual o como listado del período.",
    "Las ventas que incluyen promoción de transporte requieren un seguimiento particular. Su estado de despacho puede ser pendiente, despachado o entregado al cliente. Mientras una venta con transporte no se encuentre entregada, permanece pendiente y no descuenta stock ni se considera ingreso completado. Al confirmarse la entrega, el sistema actualiza el inventario y la operación pasa a ser una venta completada. Esta regla permite diferenciar una solicitud de transporte de una venta efectivamente entregada.",
    "Para apoyar el control administrativo, el administrador puede consultar un panel de indicadores, alertas de stock mínimo y un reporte financiero por períodos. El reporte integra ventas completadas y compras; por su parte, el vendedor solo puede consultar el detalle de las ventas que registró. La información centralizada facilita revisar resultados comerciales y tomar decisiones de abastecimiento.",
    "PEBIE dispone además de un catálogo web público donde los clientes pueden buscar artículos, filtrar por categoría, revisar imágenes, características, precios y disponibilidad. Desde el catálogo pueden solicitar una reserva indicando sus datos y la cantidad requerida. El sistema verifica la existencia, calcula el precio aplicable, genera un código de reserva y establece una fecha límite. El personal autorizado consulta y actualiza posteriormente el estado de estas reservas hasta confirmarlas, concretarlas, entregarlas o cancelarlas.",
    "En consecuencia, el problema central de Importadora PEBIE consiste en la falta de una herramienta integrada para administrar la información de usuarios, clientes, proveedores, categorías, artículos, ubicaciones, compras, inventario, mermas, kardex, ventas, despacho, reservas y reportes. Los procedimientos manuales y los registros distribuidos generan riesgo de pérdida e inconsistencia de información, retrasos en la atención y dificultades para conocer la disponibilidad real de los productos.",
    "Por esta razón, el sistema propuesto centraliza los procesos comerciales y administrativos de la empresa mediante control de acceso por roles, gestión de datos maestros, actualización automática del inventario, seguimiento de ventas con transporte, catálogo web y reservas, comprobantes imprimibles y reportes financieros. Estas funcionalidades responden al alcance implementado y a las necesidades identificadas en las historias de usuario del sistema."
]

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
normal.font.size = Pt(11)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Descripción y Dominio del Problema Actualizado')
run.font.name = 'Arial'
run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0, 0, 0)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(10)
heading.paragraph_format.space_after = Pt(10)
run = heading.add_run('3.1 DESCRIPCIÓN Y DOMINIO DEL PROBLEMA')
run.bold = True
run.font.name = 'Arial'
run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
run.font.size = Pt(12)

for text in paragraphs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.0)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    r.font.size = Pt(11)

doc.core_properties.title = 'Descripción y Dominio del Problema Actualizado PEBIE'
doc.save(OUT)
print(OUT)
